import os
import json
import sqlite3
from pathlib import Path
from decimal import Decimal
from types import SimpleNamespace
import pytest
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from portfolio_v5.models import ImportBatch,Snapshot,Position,BankTransaction
from portman.repository_v09 import RepositoryV09,reference
from portman.import_v09 import ImportSession,parse_profile_csv
from portman.analysis_v09 import change_analysis,export_report
from portman.storage_v09 import backup_records,export_package,prepare_restore,restore_prepared,unpack_package


@pytest.fixture
def repo(tmp_path):
    result=RepositoryV09(tmp_path/'portfolio.sqlite',backups=False)
    yield result
    result.close()


def job(repo): return SimpleNamespace(path=str(repo.store.path),progress=SimpleNamespace(emit=lambda *a:None),isInterruptionRequested=lambda:False)


def snapshot(day,value='100',units='1',depot='Depot',currency='EUR'):
    return ImportBatch(day+depot,day+depot+value,snapshots=[Snapshot(depot,day[:7],day,'Test',[Position('TEST','Test',quantity=units,value=value,currency=currency)])])


def booking(path='one',external='tx1',amount='100',account='Bank',day='2026-01-01'):
    return ImportBatch(path,path,transactions=[BankTransaction(account,external,day,'BANK_BOOKING',amount,'EUR',purpose='Test')],account_names=[account])


def test_revision_persists_and_open_is_read_stable(repo):
    repo.commit_batches([snapshot('2026-01-01')]); revision=repo.revision()
    other=RepositoryV09(repo.store.path,backups=False)
    try: assert other.revision()==revision
    finally: other.close()
    repo.create_accounts(['Extra']); assert repo.revision()>revision


def test_unclassified_interest_credit_is_available_to_every_overview(repo):
    batch = booking(amount='7.50')
    batch.transactions[0].purpose = 'Habenzinsen Quartal'
    repo.commit_batches([batch])

    tx = repo.transactions()[0]
    assert tx['kind'] == 'INTEREST'
    assert tx['classification_source'] == 'Automatisch aus eindeutigem Zinsbuchungstext'


def test_single_snapshot_simulations_do_not_accumulate(repo):
    session=ImportSession(repo.store.path,[snapshot('2026-01-01')])
    session.snapshot.stat().st_size
    for _ in range(3):
        result=session.simulate(job(repo),session.batches,{})
        assert result['results'][0]['status']=='imported'
    assert len(list(session.folder.glob('*.sqlite')))==1
    assert repo.holdings()==[]
    session.commit(job(repo),session.batches,{})
    assert len(repo.holdings())==1


def test_stale_import_rejected_without_partial_write(repo):
    session=ImportSession(repo.store.path,[snapshot('2026-01-01')])
    repo.create_accounts(['Changed'])
    with pytest.raises(ValueError,match='seit der Vorschau'): session.commit(job(repo),session.batches,{})
    assert repo.holdings()==[]


def test_changed_source_rejected(repo,tmp_path):
    source=tmp_path/'source.csv'; source.write_text('original')
    b=booking(str(source)); session=ImportSession(repo.store.path,[b]); source.write_text('changed')
    with pytest.raises(ValueError,match='Quelldatei'): session.commit(job(repo),[b],{})
    assert not repo.transactions()


def test_replace_is_audited_and_undo_restores_old(repo):
    repo.commit_batches([booking()]); old=repo.transactions()[0]
    b=booking('two','tx2'); session=ImportSession(repo.store.path,[b]); conflicts=session.conflicts([b]); assert len(conflicts)==1
    decisions={conflicts[0]['key']:dict(action='replace',reason='Geprüfter neuer Beleg')}
    result=session.commit(job(repo),[b],decisions)
    assert len(repo.transactions())==1 and repo.transactions()[0]['external_id'].startswith('tx2')
    assert repo.db.execute("SELECT count(*) FROM change_log WHERE entity='transaction_supersessions'").fetchone()[0]==1
    repo.undo_import(result[0]['id'])
    assert repo.transactions()==[old]


def test_keep_same_external_id_preserves_both(repo):
    repo.commit_batches([booking()]); b=booking('two'); session=ImportSession(repo.store.path,[b])
    decisions={c['key']:dict(action='keep') for c in session.conflicts([b])}
    session.commit(job(repo),[b],decisions)
    assert len(repo.transactions())==2


def test_unresolved_conflict_cannot_commit(repo):
    repo.commit_batches([booking()]); b=booking('two'); session=ImportSession(repo.store.path,[b])
    with pytest.raises(ValueError,match='ausdrücklich'): session.commit(job(repo),[b],{})
    assert len(repo.transactions())==1


def test_save_resume_session(repo):
    b=booking(); session=ImportSession(repo.store.path,[b]); saved=session.save()
    resumed=ImportSession.resume(repo.store.path,saved)
    assert resumed.batches[0].transactions[0].purpose=='Test'
    assert resumed.revision==session.revision
    resumed.commit(job(repo),resumed.batches,{})
    assert len(repo.transactions())==1


def test_atomic_multi_batch_failure(repo):
    batches=[booking(),ImportBatch('bad','bad',errors=['Broken'])]; session=ImportSession(repo.store.path,batches)
    with pytest.raises(ValueError,match='Broken'): session.commit(job(repo),batches,{})
    assert not repo.transactions()


def test_unknown_start_not_zero_and_same_coverage(repo):
    repo.create_accounts(['A','B']); repo.add_balance('A','100','2026-01-01'); repo.add_balance('A','120','2026-02-01'); repo.add_balance('B','500','2026-02-01')
    data=change_analysis(repo,'2026-01-01','2026-02-01')
    assert data['comparable_change']==Decimal(20)
    b=next(r for r in data['accounts'] if r['account']=='B'); assert b['before'] is None and b['change'] is None


def test_aggregate_units_can_be_verified(repo):
    repo.commit_batches([snapshot('2026-01-01'),snapshot('2026-02-01')]); repo.confirm_coverage('Depot','2026-01-01','2026-02-01')
    data=change_analysis(repo,'2026-01-01','2026-02-01')
    assert data['positions'][0]['status']=='Abgestimmt'


def test_missing_snapshot_is_unknown(repo):
    repo.commit_batches([snapshot('2026-02-01')])
    r=change_analysis(repo,'2026-01-01','2026-02-01')['positions'][0]
    assert r['before'] is None and r['opening_units'] is None and r['change'] is None


def test_split_replays_units_without_mutating_original(repo):
    repo.commit_batches([snapshot('2026-01-01','100','1'),snapshot('2026-02-01','100','2')])
    repo.add_action('2026-01-15','Depot','TEST','SPLIT','Beleg',ratio='2')
    row=change_analysis(repo,'2026-01-01','2026-02-01')['positions'][0]
    assert row['expected_units']==2 and row['difference']==0 and row['price_effect'] is None
    assert repo.holdings('2026-01-01')[0]['quantity']=='1'


def test_fx_decomposition_requires_evidence(repo):
    repo.commit_batches([snapshot('2026-01-01','100','1',currency='USD'),snapshot('2026-02-01','120','1',currency='USD')])
    assert change_analysis(repo,'2026-01-01','2026-02-01')['positions'][0]['fx_effect'] is None
    repo.save_fx('2026-01-01','USD','.9','Beleg'); repo.save_fx('2026-02-01','USD','.8','Beleg')
    row=change_analysis(repo,'2026-01-01','2026-02-01')['positions'][0]
    assert row['price_effect']==18 and row['fx_effect']==-12


def test_internal_transfer_scoped_cashflow(repo):
    repo.commit_batches([booking('a','a','-100','A','2026-01-10'),booking('b','b','100','B','2026-01-11')])
    a,b=repo.transactions(); repo.confirm_transfer(reference(a),reference(b),'Gegenkonto geprüft')
    assert change_analysis(repo,'2026-01-01','2026-02-01')['unclassified']==0
    assert change_analysis(repo,'2026-01-01','2026-02-01','A')['internal']['EUR']==-100
    with pytest.raises(ValueError): repo.confirm_transfer(reference(a),reference(b),'Nochmal')


def test_view_overwrite_explicit_and_undo(repo):
    repo.save_view('Test',{'a':1})
    with pytest.raises(ValueError): repo.save_view('Test',{'a':2})
    repo.save_view('Test',{'a':2},replace=True); repo.undo_view(); assert repo.saved_views()['Test']=={'a':1}
    repo.delete_view('Test'); repo.undo_view(); assert repo.saved_views()['Test']=={'a':1}


def test_manual_backup_survives_rotation(repo):
    manual=repo.backup(label='Bleibt')
    for _ in range(15): repo.backup('before-change')
    assert manual.exists()
    assert len(backup_records(repo.backup_folder))<16


def test_package_hash_and_restore(repo,tmp_path):
    repo.create_accounts(['A']); repo.add_balance('A','100','2026-01-01')
    path=tmp_path/'backup.zip'; export_package(repo.db,tmp_path,path)
    repo.add_balance('A','200','2026-02-01'); folder=tmp_path/'prepared'; prepare_restore(path,folder)
    restore_prepared(repo.db,folder,tmp_path)
    assert repo.wealth_summary()['value']==100


def test_malicious_package_path_rejected(tmp_path):
    import zipfile
    source=tmp_path/'bad.zip'
    with zipfile.ZipFile(source,'w') as archive: archive.writestr('../outside','bad')
    with pytest.raises(ValueError): unpack_package(source,tmp_path/'target')
    assert not (tmp_path/'outside').exists()


def test_csv_profile_respects_sign_and_decimal(tmp_path):
    source=tmp_path/'input.csv'; source.write_text('Account;Day;Amount;Reference\nA;01/31/2026;1,234.56;REF-77\n',encoding='utf-8')
    profile=dict(columns=dict(account='Account',date='Day',amount='Amount',bank_reference='Reference'),date_format='%m/%d/%Y',decimal='.',thousands=',',invert_sign=True)
    result=parse_profile_csv(source,profile)
    assert not result.errors and result.transactions[0].amount=='-1234.56' and result.transactions[0].date=='2026-01-31'
    assert result.transactions[0].bank_reference == result.transactions[0].document_reference == 'REF-77'


def test_analysis_invalid_range_clears_ui(repo):
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QDate
    from portman.widgets_v09 import AnalysisDialog
    QApplication.instance() or QApplication([])
    repo.commit_batches([snapshot('2026-01-01'),snapshot('2026-02-01')])
    dialog=AnalysisDialog(repo,start='2026-01-01',end='2026-02-01')
    assert dialog.tables[2].rowCount()==1
    dialog.start.setDate(QDate(2026,3,1))
    dialog._refresh_timer.timeout.emit()
    assert dialog.data is None and not dialog.export.isEnabled() and all(t.rowCount()==0 for t in dialog.tables)
    dialog.close()


def test_report_escapes_sources(repo,tmp_path):
    repo.commit_batches([snapshot('2026-01-01'),snapshot('2026-02-01')]); data=change_analysis(repo,'2026-01-01','2026-02-01'); data['scope']='<script>x</script>'
    destination=tmp_path/'report.html'; export_report(data,destination)
    assert '<script>' not in destination.read_text(encoding='utf-8') and 'Datenrevision' in destination.read_text(encoding='utf-8')


def test_fee_auxiliary_field_and_dedicated_row_count_once(repo):
    from portfolio_v5.models import Transaction
    repo.commit_batches([ImportBatch('fees','fees',transactions=[
        Transaction('A','buy','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','fee','2026-01-10','FEE','-2','EUR',security='TEST',fees='2')])])
    rows=change_analysis(repo,'2026-01-01','2026-02-01')['income']
    assert sum(r['amount'] for r in rows if r['kind']=='FEE')==-4
    a,b=repo.transactions()
    repo.confirm_fee(reference(a),'fees',reference(b),'Gleicher Beleg geprüft')
    assert sum(r['amount'] for r in change_analysis(repo,'2026-01-01','2026-02-01')['income'] if r['kind']=='FEE')==-2


def test_restoration_invalidates_previously_reviewed_session(repo,tmp_path):
    backup=repo.backup(); session=ImportSession(repo.store.path,[booking()]); folder=tmp_path/'restored'
    prepare_restore(backup,folder); restore_prepared(repo.db,folder,tmp_path)
    with pytest.raises(ValueError,match='seit der Vorschau'): session.commit(job(repo),session.batches,{})


def test_package_contains_recoverable_original_manifest(repo,tmp_path):
    source=tmp_path/'original.csv'; source.write_text('Testbeleg',encoding='utf-8'); b=booking(str(source)); repo.commit_batches([b])
    archive=tmp_path/'full.zip'; export_package(repo.db,tmp_path,archive,True)
    prepared=tmp_path/'prepared'; prepare_restore(archive,prepared); restore_prepared(repo.db,prepared,tmp_path)
    manifest=json.loads((tmp_path/'restored-sources/manifest.json').read_text(encoding='utf-8'))
    item=manifest['sources'][0]
    assert item['included'] and (tmp_path/'restored-sources'/Path(item['entry']).relative_to('sources')).read_text(encoding='utf-8')=='Testbeleg'


def test_bounded_repository_and_storage_caches(repo):
    for year in range(1900,2000): repo.holdings(f'{year}-01-01'); repo.store.holdings(f'{year}-01-01')
    assert len(repo._v8_cache)<=32 and len(repo.store._read_cache)<=64


def test_cancel_before_commit_leaves_database_unchanged(repo):
    session=ImportSession(repo.store.path,[booking()]); cancelled=job(repo); cancelled.isInterruptionRequested=lambda:True
    with pytest.raises(ValueError,match='abgebrochen'): session.commit(cancelled,session.batches,{})
    assert repo.transactions()==[]


def test_restore_file_failure_rolls_back_database(repo,tmp_path,monkeypatch):
    import portman.storage_v09 as storage
    repo.create_accounts(['A']); repo.add_balance('A','100','2026-01-01')
    prepared=tmp_path/'prepared'; prepare_restore(repo.backup(),prepared)
    (prepared/'settings').mkdir(); (prepared/'settings/portfolio_settings.json').write_text('{}')
    repo.add_balance('A','200','2026-02-01')
    original=storage.os.replace
    def fail(source,target):
        if str(source).endswith('.restore-tmp'): raise OSError('disk full')
        return original(source,target)
    monkeypatch.setattr(storage.os,'replace',fail)
    with pytest.raises(OSError,match='disk full'): restore_prepared(repo.db,prepared,tmp_path)
    assert repo.wealth_summary()['value']==200


def test_new_files_same_id_different_amount_require_review(repo):
    batches=[booking('a','same','100'),booking('b','same','200')]; session=ImportSession(repo.store.path,batches)
    conflicts=session.conflicts(batches); assert len(conflicts)==1
    with pytest.raises(ValueError,match='ausdrücklich'): session.commit(job(repo),batches,{})
    session.commit(job(repo),batches,{conflicts[0]['key']:dict(action='keep')})
    assert sorted(Decimal(t['amount']) for t in repo.transactions())==[100,200]


def test_incoming_transfer_visible_when_only_destination_selected(repo):
    def state(day,depot,qty):
        return ImportBatch(day+depot,day+depot,snapshots=[Snapshot(depot,day[:7],day,'Test',[Position('TEST','Test',quantity=str(qty),value=str(qty*10))] if qty else [])])
    repo.commit_batches([state('2026-01-01','A',10),state('2026-01-01','B',0),state('2026-02-01','A',5),state('2026-02-01','B',5)])
    repo.add_action('2026-01-15','A','TEST','TRANSFER','Original',target_depot='B',units='5')
    scoped=change_analysis(repo,'2026-01-01','2026-02-01','B')['positions']
    assert len(scoped)==1 and scoped[0]['expected_units']==5 and scoped[0]['difference']==0


def test_repetitive_conflicts_are_grouped_not_quadratic(repo):
    batch=ImportBatch('repeat','repeat',transactions=[BankTransaction('Bank',str(i),'2026-01-01','BANK_BOOKING','10','EUR') for i in range(2000)])
    session=ImportSession(repo.store.path,[batch]); conflicts=session.conflicts([batch])
    assert len(conflicts)==1999 and len(session.candidates(conflicts[-1]['key']))==1999
    assert session.conflicts([batch])==conflicts


def test_duplicate_ids_in_one_file_rejected(repo):
    b=booking(); b.transactions.append(BankTransaction('Bank','tx1','2026-01-01','BANK_BOOKING','200','EUR'))
    session=ImportSession(repo.store.path,[b])
    with pytest.raises(ValueError,match='innerhalb einer Datei'): session.conflicts([b])


def test_multiple_candidates_require_explicit_replacement_target(repo):
    repo.commit_batches([booking('a','a'),booking('b','b')]); incoming=booking('c','c'); session=ImportSession(repo.store.path,[incoming]); conflicts=session.conflicts([incoming])
    assert len(conflicts)==1
    with pytest.raises(ValueError,match='ausdrücklich auswählen'): session.commit(job(repo),[incoming],{conflicts[0]['key']:dict(action='replace',reason='Beleg')})
    target=repo.transactions()[0]
    session.commit(job(repo),[incoming],{conflicts[0]['key']:dict(action='replace',reason='Beleg',old_ref=reference(target))})
    assert len(repo.transactions())==2


def test_rebase_preserves_only_unchanged_conflicts(repo):
    repo.commit_batches([booking()]); incoming=booking('two','two'); session=ImportSession(repo.store.path,[incoming]); c=session.conflicts([incoming])[0]
    session.decisions={c['key']:dict(action='skip')}; repo.create_accounts(['Unrelated'])
    fresh=session.rebase(); assert fresh.revision==repo.revision() and fresh.decisions==session.decisions
    repo.commit_batches([booking('third','third')]); changed=fresh.rebase(); assert not changed.decisions


def test_stale_account_balance_is_not_comparable(repo):
    repo.create_accounts(['Bank']); repo.add_balance('Bank','100','2025-01-01'); repo.add_balance('Bank','200','2026-02-01')
    data=change_analysis(repo,'2026-01-01','2026-02-01')
    assert data['comparable_accounts']==0 and '45 Tage' in data['accounts'][0]['status']


def test_classification_changes_effective_not_original_booking(repo):
    repo.commit_batches([booking()]); tx=repo.transactions()[0]
    repo.classify(reference(tx),'INTEREST','',None,'Zinsbeleg')
    assert repo.transactions()[0]['kind']=='INTEREST'
    assert json.loads(repo.db.execute('SELECT payload FROM transactions').fetchone()[0])['kind']=='BANK_BOOKING'
    repo.remove_classification(reference(tx)); assert repo.transactions()[0]['kind']=='BANK_BOOKING'


def test_booking_edits_and_soft_delete_keep_original_import(repo):
    repo.commit_batches([booking()]); tx=repo.transactions()[0]; ref=reference(tx)
    repo.edit_transaction(ref, dict(date='2026-01-02', value_date='2026-01-03', kind='DIVIDEND',
        amount='12.34', currency='EUR', counterparty='Fondsgesellschaft', purpose='Ausschüttung',
        security='TEST', quantity=None), 'Kontoauszug geprüft')
    changed=repo.transactions()[0]
    assert (changed['date'],changed['kind'],changed['amount'],changed['security']) == ('2026-01-02','DIVIDEND','12.34','TEST')
    original=json.loads(repo.db.execute('SELECT payload FROM transactions').fetchone()[0])
    assert original['date']=='2026-01-01' and original['kind']=='BANK_BOOKING'
    repo.exclude_transactions([reference(changed)], 'Fehlimport')
    assert repo.transactions()==[]
    assert repo.db.execute('SELECT count(*) FROM transactions').fetchone()[0]==1


def test_duplicate_booking_is_visible_and_can_be_marked_reviewed(repo):
    first=booking('first','one'); first.transactions[0].purpose='Mitgliedsbeitrag Sportverein'
    second=booking('second','two'); second.transactions[0].purpose='Mitgliedsbeitrag Sportverein'
    repo.commit_batches([first,second])
    rows=repo.transactions()
    assert len(rows)==1 and len(rows[0]['booking_sources'])==2
    ref=reference(rows[0]); repo.resolve_duplicate(ref,'Beide Kontoauszüge geprüft')
    assert ref in repo.reviewed_duplicates()


def test_wealth_bridge_closes(repo):
    repo.create_accounts(['Bank']); repo.add_balance('Bank','100','2026-01-01'); repo.add_balance('Bank','130','2026-02-01')
    bridge=change_analysis(repo,'2026-01-01','2026-02-01')['bridge']
    assert sum(r['amount'] for r in bridge[:-1])==bridge[-1]['amount']


def test_package_encryption_authentication(tmp_path):
    from portman.package_crypto_v09 import encrypt,decrypt
    source=tmp_path/'plain'; source.write_bytes(b'private'*10000); encrypted=tmp_path/'backup.pm9'; result=tmp_path/'result'
    encrypt(source,encrypted,'long-password-123'); decrypt(encrypted,result,'long-password-123'); assert result.read_bytes()==source.read_bytes()
    result.unlink()
    with pytest.raises(ValueError,match='Passwort falsch'): decrypt(encrypted,result,'other-password-123')
    assert not result.exists()
    data=bytearray(encrypted.read_bytes()); data[-20]^=1; encrypted.write_bytes(data)
    with pytest.raises(ValueError): decrypt(encrypted,result,'long-password-123')


def test_session_cleanup_protects_resumable(repo,tmp_path):
    from portman.storage_v09 import remove_session,atomic_json
    session=ImportSession(repo.store.path,[booking()]); session.save()
    with pytest.raises(ValueError): remove_session(tmp_path,session.folder)
    atomic_json(session.folder/'discarded.json',{'at':'test'}); remove_session(tmp_path,session.folder); assert not session.folder.exists()
    with pytest.raises(ValueError): remove_session(tmp_path,tmp_path)


def test_startup_recovers_interrupted_restore(repo,tmp_path):
    from portman.storage_v09 import recover_pending,atomic_json
    repo.create_accounts(['Bank']); repo.add_balance('Bank','100','2026-01-01')
    stage=tmp_path/'restore-jobs'/'test'; stage.mkdir(parents=True)
    dst=sqlite3.connect(stage/'rollback.sqlite'); repo.db.backup(dst); dst.close()
    original=tmp_path/'portfolio_settings.json'; original.write_text('{"old":1}'); (stage/'0').write_text('{"old":1}')
    atomic_json(stage/'journal.json',dict(status='pending',files=[dict(target='portfolio_settings.json',previous='0')]))
    repo.add_balance('Bank','999','2026-02-01'); original.write_text('{"new":2}')
    assert recover_pending(tmp_path,repo.db)==['test']
    assert repo.wealth_summary()['value']==100 and json.loads(original.read_text())=={'old':1}
    assert not recover_pending(tmp_path,repo.db)


def test_excel_export_has_german_headers_and_unknowns(repo,tmp_path):
    from portman.reports_v09 import export_report
    from openpyxl import load_workbook
    repo.commit_batches([snapshot('2026-02-01')]); data=change_analysis(repo,'2026-01-01','2026-02-01'); path=tmp_path/'report.xlsx'; export_report(data,path)
    wb=load_workbook(path); assert wb['Positionen']['A1'].value=='Konto' and wb['Positionen']['E2'].value=='Unbekannt'; assert wb['Positionen'].freeze_panes=='C2'; wb.close()


def test_virtual_table_sorts_numeric_and_fetches_more():
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import Qt
    from portman.tables_v09 import RowsModel
    QApplication.instance() or QApplication([]); model=RowsModel(['Wert'],[[Decimal(i)] for i in range(1000)])
    assert model.rowCount()==256; model.sort(0,Qt.SortOrder.DescendingOrder); assert model.data(model.index(0,0),Qt.ItemDataRole.UserRole)==999
    model.fetchMore(); assert model.rowCount()==512; model.filter('999'); assert len(model.rows)==1


def test_fee_cannot_link_to_itself(repo):
    from portfolio_v5.models import Transaction
    repo.commit_batches([ImportBatch('fee','fee',transactions=[Transaction('A','fee','2026-01-10','FEE','-2','EUR',fees='2')])])
    ref=reference(repo.transactions()[0])
    with pytest.raises(ValueError,match='unterschiedliche'): repo.confirm_fee(ref,'fees',ref,'Test')
    assert repo.fee_links()==[]


def test_action_correction_is_atomic(repo,monkeypatch):
    values=dict(day='2026-01-10',depot='A',security='TEST',kind='SPLIT',source='Beleg',ratio='2')
    repo.add_action(**values); original=repo.actions('2026-01-01','2026-02-01')[0]
    def fail(*args): raise RuntimeError('Simulierter Schreibfehler')
    with monkeypatch.context() as patch:
        patch.setattr(repo,'revoke_action',fail)
        with pytest.raises(RuntimeError): repo.correct_action(original['id'],**{**values,'ratio':'3'})
    assert repo.actions('2026-01-01','2026-02-01')==[original]
    repo.correct_action(original['id'],**{**values,'ratio':'3'})
    assert len(repo.actions('2026-01-01','2026-02-01'))==1
    assert repo.actions('2026-01-01','2026-02-01')[0]['ratio']=='3'


def test_completed_read_job_clears_running_status():
    from Gui.PortMan_V09 import MainWindow
    from PySide6.QtWidgets import QApplication,QLabel
    QApplication.instance() or QApplication([])
    label=QLabel('Verarbeitung läuft …')
    window=SimpleNamespace(repository=SimpleNamespace(poll_revision=lambda:7),job_center=SimpleNamespace(records=[{'end':1}]),saved_status=label,_last_revision=7)
    MainWindow.update_save_status(window)
    assert label.text().startswith('Bereit')
    label.setText('Speichern fehlgeschlagen')
    MainWindow.update_save_status(window)
    assert label.text()=='Speichern fehlgeschlagen'


def test_dashboard_defers_initial_queries():
    from portman.widgets_v09 import DashboardV09
    DashboardV09.refresh(SimpleNamespace(repo=SimpleNamespace(defer_initial_views=True)))


def test_edit_dialog_refreshes_only_after_changes():
    from Gui.PortMan_V09 import MainWindow
    state={'revision':1,'reloads':0}
    def reload(): state['reloads']+=1
    window=SimpleNamespace(repository=SimpleNamespace(revision=lambda:state['revision']),_scope_dirty=set(),_scope_refreshers={'dashboard':None},reload_portfolio=reload)
    MainWindow.run_edit_dialog(window,SimpleNamespace(exec=lambda:0))
    assert state['reloads']==0
    def edit(): state['revision']+=1; return 1
    assert MainWindow.run_edit_dialog(window,SimpleNamespace(exec=edit))==1
    assert state['reloads']==1 and window._scope_dirty=={'dashboard'}


def test_transfer_link_no_longer_hides_reclassified_income(repo):
    repo.commit_batches([booking('a','a','-100','A','2026-01-10'),booking('b','b','100','B','2026-01-11')])
    a,b=repo.transactions(); repo.confirm_transfer(reference(a),reference(b),'Geprüft')
    repo.classify(reference(b),'INTEREST','','','Zinsbeleg')
    assert repo.transfer_pairs()==[]
    stored=repo.transfer_pairs(include_invalid=True)
    assert len(stored)==1
    assert change_analysis(repo,'2026-01-01','2026-02-01')['explained']['INTEREST']==100
    repo.revoke_transfer(stored[0]['id'],'Korrigierte Zuordnung')
    assert repo.transfer_pairs(include_invalid=True)==[]


def test_transfer_confirmation_rejects_income(repo):
    repo.commit_batches([booking('a','a','-100','A'),booking('b','b','100','B')])
    a,b=repo.transactions(); repo.classify(reference(b),'DIVIDEND','','','Beleg')
    with pytest.raises(ValueError): repo.confirm_transfer(reference(a),reference(b),'Kein Transfer')


def test_fee_candidates_exclude_consumed_counterpart(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import fee_candidates
    repo.commit_batches([ImportBatch('fees','fees',transactions=[
        Transaction('A','buy1','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','buy2','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','fee','2026-01-10','FEE','-2','EUR')])])
    a,b,c=repo.transactions(); repo.confirm_fee(reference(a),'fees',reference(c),'Geprüft')
    assert fee_candidates(repo)==[]


def test_fee_candidates_match_with_three_day_valuta_skew(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import fee_candidates
    repo.commit_batches([ImportBatch('fees','fees',transactions=[
        Transaction('A','buy1','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','fee','2026-01-12','FEE','-2','EUR')])])
    candidates = fee_candidates(repo)
    assert [(c['field'], c['amount']) for c in candidates] == [('fees', '2')]


def test_fee_candidates_never_bridge_more_than_three_days(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import fee_candidates
    repo.commit_batches([ImportBatch('fees','fees',transactions=[
        Transaction('A','buy1','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','fee','2026-01-15','FEE','-2','EUR')])])
    assert fee_candidates(repo)==[]


def test_revoked_fee_link_reopens_the_candidate(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import fee_candidates
    repo.commit_batches([ImportBatch('fees','fees',transactions=[
        Transaction('A','buy1','2026-01-10','BUY','-100','EUR',security='TEST',quantity='1',fees='2'),
        Transaction('A','fee','2026-01-10','FEE','-2','EUR')])])
    a,b=repo.transactions(); repo.confirm_fee(reference(a),'fees',reference(b),'Geprüft')
    assert fee_candidates(repo)==[]
    # An invalidated link no longer hides the candidate: the analysis counts
    # the fee twice again, so the proposal must come back.
    repo.revoke_fee(reference(a),'fees')
    assert len(fee_candidates(repo))==1


def test_recurring_equal_amounts_with_different_decimal_precision(repo):
    from portman.analysis_v09 import recurring_candidates
    repo.commit_batches([booking('a','a','10','A','2026-01-01'),booking('b','b','10.0','A','2026-02-01'),booking('c','c','10.00','A','2026-03-01')])
    assert len(recurring_candidates(repo))==1


def test_fifo_target_of_merger_is_not_reported_as_known(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import realized_profit
    repo.commit_batches([ImportBatch('trades','trades',transactions=[
        Transaction('A','buy','2026-01-01','BUY','-100','EUR',security='NEW',quantity='1'),
        Transaction('A','sell','2026-02-01','SELL','150','EUR',security='NEW',quantity='1')])])
    repo.confirm_coverage('A','0001-01-01','2026-03-01')
    repo.add_action('2026-01-15','A','OLD','MERGER','Beleg',target='NEW')
    assert realized_profit(repo,'2026-01-01','2026-03-01')[0]['realized'] is None


def test_duplicate_ids_with_equivalent_account_names_rejected(repo):
    batch=booking(account='ING')
    batch.transactions+=booking('two',account='Ing Diba',amount='200').transactions
    session=ImportSession(repo.store.path,[batch])
    with pytest.raises(ValueError,match='Buchungskennung'): session.conflicts([batch])


def test_one_counterpart_in_two_indexes_does_not_require_selection(repo):
    repo.commit_batches([booking()])
    batch=booking('updated'); session=ImportSession(repo.store.path,[batch])
    conflict=session.conflicts([batch])[0]
    selected,replacements=session.resolve([batch],{conflict['key']:{'action':'replace','reason':'Korrigierter Beleg'}})
    assert len(replacements)==1


def test_restore_diff_includes_import_activation_and_coverage(repo,tmp_path):
    from portman.storage_v09 import content_diff
    repo.commit_batches([booking()]); backup=repo.backup()
    with repo.db: repo.db.execute('UPDATE imports SET active=0')
    repo.confirm_coverage('Bank','2026-01-01','2026-02-01')
    diff=content_diff(repo.store.path,backup)
    assert {'imports','coverage'} <= {row[0] for row in diff['rows']}


def test_restore_diff_summarizes_binary_documents(repo):
    from base64 import b64encode
    from portfolio_v5.models import BankDocument
    from portman.storage_v09 import content_diff
    backup=repo.backup()
    payload=b'%PDF-test\n%%EOF'
    results=repo.commit_batches([ImportBatch('document','document',documents=[BankDocument('Bank','doc','Beleg','2026-01-01',b64encode(payload).decode())])])
    assert results[0]['status']=='imported'
    diff=content_diff(repo.store.path,backup)
    document=next(r for r in diff['rows'] if r[0]=='bank_documents')
    content=json.loads(document[3])['content']
    assert content['bytes']==len(payload) and len(content['sha256'])==64


@pytest.mark.parametrize('record',[dict(target='../outside.json',previous=None),dict(target='portfolio_settings.json',previous='missing'),dict(target='portfolio_settings.json',previous='../outside')])
def test_recovery_validates_files_before_replacing_database(repo,tmp_path,record):
    from portman.storage_v09 import recover_pending,atomic_json
    repo.create_accounts(['Bank']); repo.add_balance('Bank','100','2026-01-01')
    stage=tmp_path/'restore-jobs'/'invalid'; stage.mkdir(parents=True)
    dst=sqlite3.connect(stage/'rollback.sqlite'); repo.db.backup(dst); dst.close()
    repo.add_balance('Bank','999','2026-02-01')
    untouched=tmp_path/'etf_cache.json'; untouched.write_text('{}')
    atomic_json(stage/'journal.json',dict(status='pending',files=[dict(target='etf_cache.json',previous=None),record]))
    with pytest.raises(ValueError): recover_pending(tmp_path,repo.db)
    assert repo.wealth_summary()['value']==999
    assert untouched.read_text()=='{}'


def test_sale_without_quantity_remains_visible_as_unknown(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import realized_profit
    repo.commit_batches([ImportBatch('sale','sale',transactions=[Transaction('A','sell','2026-01-10','SELL','150','EUR',security='TEST')])])
    rows=realized_profit(repo,'2026-01-01','2026-02-01')
    assert len(rows)==1 and rows[0]['proceeds']==150 and rows[0]['realized'] is None


def test_advance_tax_reduces_only_later_taxable_sale_gain(repo):
    from portfolio_v5.models import Transaction, TaxEvent
    from portman.analysis_v09 import realized_profit
    repo.save_account('Depot A', {'kind': 'Depot'})
    repo.commit_batches([ImportBatch('trades','trades',transactions=[
        Transaction('Depot A','buy','2025-01-01','BUY','-200','EUR',security='IE00BKM4GZ66',quantity='2'),
        Transaction('Depot A','sell-1','2026-02-01','SELL','150','EUR',security='IE00BKM4GZ66',quantity='1'),
        Transaction('Depot A','sell-2','2026-03-01','SELL','150','EUR',security='IE00BKM4GZ66',quantity='1')])])
    repo.commit_batches([ImportBatch('tax.pdf','tax',tax_events=[TaxEvent(
        'advance','Depot A','vorabpauschale_ausland','2026-01-02',income_year='2025',
        security_ref='IE00BKM4GZ66',quantity='2',advance_tax_base='85.72',taxable_income='60.00')])])
    repo.confirm_coverage('Depot A','0001-01-01','2026-04-01')
    rows=realized_profit(repo,'2026-01-01','2026-04-01','Depot A')
    assert [(r['cost'],r['realized'],r['advance_tax_credit'],r['taxable_realized']) for r in rows] == [
        (Decimal('100'),Decimal('50'),Decimal('30.00'),Decimal('20.00')),
        (Decimal('100'),Decimal('50'),Decimal('30.00'),Decimal('20.00'))]
    assert repo.advance_tax_credit('Depot A','IE00BKM4GZ66') == Decimal('60.00')


def test_sale_before_report_period_consumes_advance_tax_credit(repo):
    from portfolio_v5.models import Transaction, TaxEvent
    from portman.analysis_v09 import realized_profit
    repo.save_account('Depot A', {'kind': 'Depot'})
    repo.commit_batches([ImportBatch('all','all',transactions=[
        Transaction('Depot A','buy','2025-01-01','BUY','-200','EUR',security='FUND',quantity='2'),
        Transaction('Depot A','old-sale','2026-02-01','SELL','150','EUR',security='FUND',quantity='1'),
        Transaction('Depot A','period-sale','2026-03-01','SELL','150','EUR',security='FUND',quantity='1')],
        tax_events=[TaxEvent('advance','Depot A','vorabpauschale','2026-01-02',security_ref='FUND',
            quantity='2',taxable_income='60')])])
    repo.confirm_coverage('Depot A','0001-01-01','2026-04-01')
    row=realized_profit(repo,'2026-02-15','2026-04-01','Depot A')[0]
    assert row['advance_tax_credit'] == Decimal('30')


def test_later_corporate_action_does_not_hide_earlier_realized_profit(repo):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import realized_profit
    repo.commit_batches([ImportBatch('trades','trades',transactions=[
        Transaction('A','buy','2026-01-02','BUY','-200','EUR',security='TEST',quantity='2'),
        Transaction('A','sell','2026-01-10','SELL','150','EUR',security='TEST',quantity='1'),
        Transaction('A','sell-after','2026-02-10','SELL','150','EUR',security='TEST',quantity='1')])])
    repo.confirm_coverage('A','0001-01-01','2026-03-01')
    repo.add_action('2026-02-01','A','TEST','SPLIT','Späterer Split',ratio='2')
    rows=realized_profit(repo,'2026-01-01','2026-03-01')
    assert rows[0]['realized']==50 and rows[1]['realized'] is None


@pytest.mark.parametrize('ids',[('a-buy','z-sell'),('z-buy','a-sell')])
def test_same_day_trade_order_is_not_inferred_from_external_ids(repo,ids):
    from portfolio_v5.models import Transaction
    from portman.analysis_v09 import realized_profit
    repo.commit_batches([ImportBatch('trades','trades',transactions=[
        Transaction('A',ids[0],'2026-01-10','BUY','-100','EUR',security='TEST',quantity='1'),
        Transaction('A',ids[1],'2026-01-10','SELL','150','EUR',security='TEST',quantity='1')])])
    repo.confirm_coverage('A','0001-01-01','2026-02-01')
    assert realized_profit(repo,'2026-01-01','2026-02-01')[0]['realized'] is None


def test_confirmed_recurring_rule_is_not_suggested_again(repo):
    from portman.analysis_v09 import recurring_candidates
    repo.commit_batches([booking('a','a','10','A','2026-01-01'),booking('b','b','10','A','2026-02-01'),booking('c','c','10','A','2026-03-01')])
    repo.confirm_recurring(recurring_candidates(repo)[0])
    assert recurring_candidates(repo)==[]
    with repo.db: repo.db.execute('DELETE FROM recurring_rules')
    assert len(recurring_candidates(repo))==1
