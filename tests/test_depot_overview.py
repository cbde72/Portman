import csv
import os
from decimal import Decimal

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtWidgets import QApplication, QFileDialog

from portfolio_v5.models import BankTransaction, ImportBatch, Position, Snapshot, TaxEvent
from portman.depot_overview import aggregate_depot
from portman.gui.depot_overview import COLUMNS, DepotOverviewWidget
from portman.repository_v09 import RepositoryV09


def build_repo(tmp_path):
    repo = RepositoryV09(tmp_path / 'overview.sqlite', backups=False)
    repo.save_account('Depot A', {'kind': 'Depot', 'institution': 'Test'})
    repo.save_account('Depot B', {'kind': 'Depot', 'institution': 'Test'})
    batch = ImportBatch('activity.pdf', 'activity-hash', account_names=['Depot A'],
        snapshots=[Snapshot('Depot A', '2026-12', '2026-12-31', 'Test', [
            Position('SEC-1', 'Alpha Fonds', isin='DE0000000001', quantity='5', value='70', cost='51')])],
        transactions=[
            BankTransaction('Depot A', 'buy-1', '2026-01-10', 'BUY', '-102', 'EUR',
                security='DE0000000001', quantity='10', fees='2', gross_amount='100'),
            BankTransaction('Depot A', 'sell-1', '2026-06-10', 'SELL', '58', 'EUR',
                security='DE0000000001', quantity='5', fees='2', gross_amount='60'),
            BankTransaction('Depot A', 'div-1', '2026-07-10', 'DIVIDEND', '8', 'EUR',
                security='DE0000000001', taxes='2', gross_amount='10',
                solidarity_tax='.1', withholding_tax='.5'),
            BankTransaction('Depot A', 'fee-1', '2026-08-10', 'FEE', '-3', 'EUR', fees='3'),
            BankTransaction('Depot B', 'other-depot', '2026-02-10', 'BUY', '-20', 'EUR',
                security='DE0000000002', quantity='1'),
        ])
    assert repo.commit_batches([batch])[0]['status'] == 'imported'
    repo.confirm_coverage('Depot A', '0001-01-01', '2026-12-31')
    return repo, batch


def test_aggregate_single_depot_period_fees_taxes_dividend_and_fifo(tmp_path):
    repo, _batch = build_repo(tmp_path)
    try:
        rows = aggregate_depot(repo, 'Depot A', '2026-01-01', '2026-12-31')
        alpha = next(row for row in rows if row['identifier'] == 'DE0000000001')
        assert alpha['quantity_end'] == Decimal('5')
        assert (alpha['buy_count'], alpha['buy_total'], alpha['buy_fees']) == (1, Decimal('100'), Decimal('2'))
        assert (alpha['sell_count'], alpha['sell_total'], alpha['sell_fees']) == (1, Decimal('60'), Decimal('2'))
        assert (alpha['dividend_gross'], alpha['dividend_net']) == (Decimal('10'), Decimal('8'))
        assert alpha['tax_total'] == Decimal('2')
        assert alpha['solidarity_tax'] == Decimal('.1')
        assert alpha['withholding_tax'] == Decimal('.5')
        assert alpha['capital_gains_tax'] == Decimal('1.4')
        assert alpha['net_cashflow'] == Decimal('-36')
        assert alpha['realized_profit'] == Decimal('7')
        unassigned = next(row for row in rows if row['key'] == '__unassigned__')
        assert unassigned['other_fees'] == Decimal('3')
        assert unassigned['net_cashflow'] == Decimal('-3')
        assert all(row['identifier'] != 'DE0000000002' for row in rows)
    finally:
        repo.close()


def test_period_and_closed_position_filter_source(tmp_path):
    repo, _batch = build_repo(tmp_path)
    try:
        before_sale = aggregate_depot(repo, 'Depot A', '2026-01-01', '2026-05-31')
        alpha = next(row for row in before_sale if row['identifier'] == 'DE0000000001')
        assert alpha['sell_count'] == 0 and alpha['dividend_net'] == 0
        # A security traded in the period remains represented even without an end holding.
        all_rows = aggregate_depot(repo, 'Depot B', '2026-01-01', '2026-12-31')
        assert len(all_rows) == 1 and all_rows[0]['buy_count'] == 1
    finally:
        repo.close()


def test_reimport_does_not_double_count(tmp_path):
    repo, batch = build_repo(tmp_path)
    try:
        assert repo.commit_batches([batch])[0]['status'] == 'duplicate'
        alpha = next(row for row in aggregate_depot(repo, 'Depot A', '2026-01-01', '2026-12-31')
                     if row['identifier'] == 'DE0000000001')
        assert alpha['buy_count'] == alpha['sell_count'] == 1
    finally:
        repo.close()


def test_reinvested_distribution_exemption_and_advance_tax_are_differentiated(tmp_path):
    repo, _batch = build_repo(tmp_path)
    try:
        extra = ImportBatch('tax-notices.pdf', 'tax-notices', tax_events=[
            TaxEvent('advance-tax', 'Depot A', 'vorabpauschale', '2026-01-04',
                     income_year='2025', security_ref='DE0000000001',
                     advance_tax_base='5', paid_taxes='1'),
            TaxEvent('exempt-dividend', 'Depot A', 'dividende_inland', '2026-09-01',
                     income_year='2026', security_ref='DE0000000001',
                     gross_income='4', used_exemption='4', paid_taxes='0'),
        ], transactions=[
            BankTransaction('Depot A', 'reinvested-income', '2026-09-01', 'DISTRIBUTION',
                            '0', 'EUR', security='DE0000000001', gross_amount='4'),
        ])
        assert repo.commit_batches([extra])[0]['status'] == 'imported'
        alpha = next(row for row in aggregate_depot(repo, 'Depot A', '2026-01-01', '2026-12-31')
                     if row['identifier'] == 'DE0000000001')
        assert alpha['dividend_gross'] == Decimal('14')
        assert alpha['dividend_net'] == Decimal('8')
        assert alpha['tax_total'] == Decimal('3')
        # The documented advance-tax burden is not an additional cash booking.
        assert alpha['net_cashflow'] == Decimal('-36')
    finally:
        repo.close()


def test_widget_columns_filters_copy_and_auto_refresh(tmp_path):
    app = QApplication.instance() or QApplication([])
    repo, _batch = build_repo(tmp_path)
    try:
        widget = DepotOverviewWidget(repo)
        widget.depot.setCurrentIndex(widget.depot.findData('Depot A'))
        assert widget.timer.isActive()
        widget.timer.stop(); widget.start.setDate(widget.start.minimumDate().addYears(126)); widget.end.setDate(widget.end.maximumDate())
        widget.refresh()
        assert widget.table.columnCount() == len(COLUMNS)
        assert widget.table.isColumnHidden(9)
        assert widget.table.rowCount() >= 2  # position plus sum
        widget.copy_all()
        assert 'Wertpapiername' in app.clipboard().text()
        widget.positions.setCurrentText('Aktive Positionen')
        assert widget.table.rowCount() >= 2
        widget.close()
    finally:
        repo.close()


def test_csv_and_xlsx_export_include_hidden_columns(tmp_path, monkeypatch):
    QApplication.instance() or QApplication([])
    repo, _batch = build_repo(tmp_path)
    try:
        widget = DepotOverviewWidget(repo); widget.depot.setCurrentIndex(widget.depot.findData('Depot A'))
        widget.start.setDate(QApplication.instance().property('unused') or widget.start.date().addYears(-10)); widget.refresh()
        csv_path = tmp_path / 'overview.csv'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(csv_path), 'CSV (*.csv)'))
        monkeypatch.setattr('portman.gui.depot_overview.QMessageBox.information', lambda *args: None)
        widget.export()
        with csv_path.open(encoding='utf-8-sig') as stream:
            header = next(csv.reader(stream, delimiter=';'))
        assert 'Kirchensteuer' in header and len(header) == len(COLUMNS)
        xlsx_path = tmp_path / 'overview.xlsx'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(xlsx_path), 'Excel (*.xlsx)'))
        widget.export()
        from openpyxl import load_workbook
        assert load_workbook(xlsx_path, read_only=True).active.max_column == len(COLUMNS)
        widget.close()
    finally:
        repo.close()
