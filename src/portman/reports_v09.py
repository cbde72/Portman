"""German, source-aware HTML and structured spreadsheet exports."""
from pathlib import Path
from html import escape
from decimal import Decimal
import json

SECTIONS=[('Konten','accounts',[('account','Konto'),('before','Anfang EUR'),('after','Ende EUR'),('change','Vergleichbare Änderung EUR'),('status','Vergleichbarkeit'),('dates_before','Anfangsbelege'),('dates_after','Endbelege')]),
 ('Positionen','positions',[('account','Konto'),('depot','Depot'),('name','Wertpapier'),('currency','Währung'),('before','Anfangswert'),('after','Endwert'),('change','Wertänderung'),('book_profit_change','Unrealisierte Gewinnänderung'),('price_effect','Kursbeitrag EUR'),('fx_effect','Währungsbeitrag EUR'),('status','Nachweisstatus'),('source_before','Anfangsquelle'),('source_after','Endquelle')]),
 ('Vermögensbrücke','bridge',[('label','Bestandteil'),('amount','EUR')]),
 ('Erträge und Kosten','income',[('month','Monat'),('account','Konto'),('security','Wertpapier'),('kind','Buchungsart'),('currency','Währung'),('amount','Betrag')]),
 ('Realisierte Gewinne','profit',[('date','Datum'),('account','Konto'),('security','Wertpapier'),('currency','Währung'),('proceeds','Verkaufserlös'),('cost','FIFO-Einstand'),('realized','Realisierter Gewinn'),('advance_tax_credit','Angerechnete Vorabpauschale'),('taxable_realized','Steuerlicher Gewinn nach Vorabpauschale'),('status','Methode / Abdeckung'),('source','Quelle')]),
 ('Wechselkursbelege','fx_sources',[('day','Datum'),('currency','Währung'),('rate','EUR je Einheit'),('source','Quelle')])]
METHOD='Dokumentierte Vermögensänderung, keine Rendite. Kurs-/FX-Näherung auf Anfangsstückzahlen; weitere Mengen- und Bewertungseffekte verbleiben im Rest. Historische FX-Belege maximal sieben Tage alt. Realisierte Gewinne: FIFO nur bei belegtem Anfangsbestand und vollständiger Buchungsabdeckung. Bereits besteuerte Vorabpauschalen werden stückanteilig ausschließlich vom steuerlichen Veräußerungsgewinn abgezogen; Einstand und Performance bleiben unverändert. Buchungsbeträge folgen der Gebührenkonvention der Quelle. Gebühren werden nur mit bestätigter Belegverknüpfung zusammengeführt. Fremdwährungen werden nicht ungeprüft als EUR summiert.'
def printable(value):
    if value is None: return 'Unbekannt'
    if isinstance(value,(dict,list)): return json.dumps(value,ensure_ascii=False,default=str)
    return str(value)
def export_report(data,destination):
    destination=Path(destination)
    meta=[('Auswahl',data['scope']),('Von',data['start']),('Bis',data['end']),('Datenrevision',data['revision']),('Datenabdeckung',printable(data['quality'])),('Methode',METHOD)]
    if destination.suffix.lower()=='.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook(); ws=wb.active; ws.title='Auswahl und Methode'
        for k,v in meta: ws.append([k,printable(v)])
        ws.column_dimensions['A'].width=25; ws.column_dimensions['B'].width=100
        for row in ws: row[1].alignment=Alignment(wrap_text=True,vertical='top')
        for title,key,columns in SECTIONS:
            sheet=wb.create_sheet(title[:31]); sheet.append([label for _,label in columns])
            for row in data[key]:
                values=[]
                for field,label in columns:
                    value=row.get(field)
                    if isinstance(value,Decimal): value=float(value)
                    elif value is None: value='Unbekannt'
                    elif not isinstance(value,(int,float)): value=printable(value)
                    if isinstance(value,str) and value.startswith(('=','+','-','@')): value="'"+value
                    values.append(value)
                sheet.append(values)
            sheet.freeze_panes='C2'; sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='255FD6')
            for c in range(1,len(columns)+1): sheet.column_dimensions[get_column_letter(c)].width=24
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type=='n': cell.number_format='#,##0.00;[Red]-#,##0.00'
        wb.save(destination); return
    html='<!doctype html><meta charset="utf-8"><title>PortMan 09 · Auswertung</title><style>body{font:14px system-ui;margin:32px}table{border-collapse:collapse}td,th{padding:6px;border:1px solid #bbb;text-align:left}th{background:#eaf0ff}p{max-width:1000px}</style><h1>PortMan 09 · Auswertung</h1>'
    html+=''.join('<p><b>'+escape(k)+':</b> '+escape(printable(v))+'</p>' for k,v in meta)
    for title,key,columns in SECTIONS:
        html+='<h2>'+title+'</h2><table><tr>'+''.join('<th>'+escape(label)+'</th>' for _,label in columns)+'</tr>'
        html+=''.join('<tr>'+''.join('<td>'+escape(printable(row.get(field)))+'</td>' for field,_ in columns)+'</tr>' for row in data[key])+'</table>'
    destination.write_text(html,encoding='utf-8')
